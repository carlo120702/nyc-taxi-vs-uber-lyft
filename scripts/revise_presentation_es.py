"""Copy NYC_Taxi_Amarillo_vs_Uber_Lyft_2025_ES.pptx and apply content fixes."""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "NYC_Taxi_Amarillo_vs_Uber_Lyft_2025_ES.pptx"
DST = ROOT / "NYC_Taxi_Amarillo_vs_Uber_Lyft_2025_ES_REVISADO.pptx"
DST_REPO = Path(__file__).resolve().parents[1] / "NYC_Taxi_Amarillo_vs_Uber_Lyft_2025_ES_REVISADO.pptx"

REPLACEMENTS = [
    (
        "287 millones de viajes, Nueva York 2025",
        "287,1 M viajes limpios · ~292 M crudos · NYC 2025",
    ),
    (
        "detectar meses nuevos → datos de contexto → limpiar con Dask → 8 reglas de calidad → agregar con Dask → cargar",
        "discover → contexto (zonas, clima) → armonizar esquemas y limpiar (Dask) → validar (10 reglas/mes) → agregar KPI (Dask) → cargar RDS → log ETL",
    ),
    (
        "PostgreSQL, solo tablas de KPI, registro de calidad y de ejecuciones",
        "PostgreSQL: 4 tablas KPI + quality_check + etl_run_log. Viajes crudos nunca salen del equipo local",
    ),
    (
        "También son vacíos estructurales, no errores: las propinas en efectivo del taxi no se registran y el taxímetro no tiene hora de solicitud ni pago al conductor. Estos vacíos se documentan, no se imputan.",
        "Yellow Taxi y FHVHV tienen esquemas distintos (p. ej. fare_amount vs. base_passenger_fare); el pipeline los armoniza antes de limpiar. "
        "También hay vacíos estructurales: propinas en efectivo del taxi no se registran; no hay hora de solicitud ni pago al conductor. Se documentan, no se imputan.",
    ),
    (
        "0,867\n---\nROC-AUC, gradient boosting sobre 500 k viajes. PR-AUC 0,51 vs. 0,15 al azar.",
        "0,867\n---\nROC-AUC (HistGradientBoosting, 500 k viajes). PR-AUC 0,51 vs. 0,15 al azar.",
    ),
    (
        "ROC-AUC, gradient boosting sobre 500 k viajes. PR-AUC 0,51 vs. 0,15 al azar.",
        "ROC-AUC (HistGradientBoosting, 500 k viajes). PR-AUC 0,51 vs. 0,15 al azar.",
    ),
    (
        "15 paneles sobre 4 tablas de AWS",
        "15 paneles · 6 tablas en RDS (4 KPI + calidad + log ETL)",
    ),
    (
        "Solo lugar y hora logran un-ROC-AUC de 0,87. El clima no cambia la cuota.",
        "Solo lugar y hora logran ROC-AUC 0,87. El clima no cambia la cuota de mercado.",
    ),
    (
        "Siguiente: modelo completo con variables a nivel de zona, añadir 2024 para ver la tendencia y programar el flujo de Prefect cada mes.",
        "Siguiente: modelo completo con variables de zona, añadir 2024 para tendencia y programar Prefect mensualmente. "
        "Flujo incremental: segunda corrida sin archivos nuevos no reprocesa; un mes nuevo solo actualiza ese mes.",
    ),
    (
        "Código: github.com/carlo120702/nyc-taxi-vs-uber-lyft · 3 / 10",
        "Código: github.com/carlo120702/nyc-taxi-vs-uber-lyft · ETL: flows/etl_flow.py · 3 / 10",
    ),
]


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def replace_in_text_frame(text_frame, replacements):
    changed = False
    for para in text_frame.paragraphs:
        for run in para.runs:
            original = run.text
            new = original
            for old, repl in replacements:
                if old in new:
                    new = new.replace(old, repl)
            if new != original:
                run.text = new
                changed = True
    return changed


def replace_whole_shape(shape, replacements) -> bool:
    """When text is split across runs, replace using the full shape text."""
    if not shape.has_text_frame:
        return False
    text = shape.text
    new_text = text
    for old, repl in replacements:
        if old in new_text:
            new_text = new_text.replace(old, repl)
    if new_text == text:
        return False
    tf = shape.text_frame
    first = True
    for para in tf.paragraphs:
        if first:
            if para.runs:
                para.runs[0].text = new_text
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.text = new_text
            first = False
        else:
            para.text = ""
            for run in para.runs:
                run.text = ""
    return True


def apply_replacements(prs: Presentation) -> int:
    n = 0
    for slide in prs.slides:
        for shape in iter_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            if replace_in_text_frame(shape.text_frame, REPLACEMENTS):
                n += 1
            elif replace_whole_shape(shape, REPLACEMENTS):
                n += 1
    return n


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")
    prs = Presentation(str(SRC))
    count = apply_replacements(prs)
    prs.save(str(DST))
    prs.save(str(DST_REPO))
    print(f"Saved: {DST}")
    print(f"Saved: {DST_REPO}")
    print(f"Shapes updated: {count}")


if __name__ == "__main__":
    main()
